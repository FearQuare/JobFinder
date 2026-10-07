"""The user's profile: what they know, what they want. Stored locally as one JSON document."""
from __future__ import annotations

import io
import json
import re
import sqlite3
from typing import Any, Dict, List, Optional, Tuple

from . import ai, config, db

SKILL_LEVELS = {1: "Beginner", 2: "Intermediate", 3: "Advanced", 4: "Expert"}
LANGUAGE_LEVELS = ["A1", "A2", "B1", "B2", "C1", "C2", "native"]
EDU_STATUS = ["completed", "in_progress", "planned"]
MAX_CV_BYTES = 8 * 1024 * 1024
CV_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")

EMPTY_PROFILE: Dict[str, Any] = {
    "name": "", "headline": "", "location": "", "summary": "",
    "links": {"github": "", "linkedin": "", "portfolio": ""},
    "target_categories": [], "target_seniorities": [], "target_job_types": [], "preferred_locations": "",
    "weekly_hours": 10, "timeline_months": 6,
    "education": [], "experience": [], "skills": [], "languages": [], "certifications": [], "projects": [],
}

PROFILE_SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS profile (
    id INTEGER PRIMARY KEY CHECK (id = 1),
    data TEXT NOT NULL,
    cv_filename TEXT,
    cv_text TEXT,
    updated_at TEXT NOT NULL
);
"""


# ---------------------------------------------------------------- cleaning

def _s(v: Any, limit: int = 300) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()[:limit]


def _text(v: Any, limit: int = 3000) -> str:
    return str(v or "").strip()[:limit]


def _int(v: Any, lo: int, hi: int, default: int) -> int:
    try:
        return max(lo, min(hi, int(float(v))))
    except (TypeError, ValueError):
        return default


def _list(v: Any) -> List[str]:
    if isinstance(v, str):
        v = v.split(",")
    return [x for x in (_s(i, 60) for i in (v or [])) if x]


def language_level(level: Any) -> str:
    lvl = _s(level, 60).upper()
    if "NATIVE" in lvl or "MUTTER" in lvl or lvl in ("C2+", "NATIVE SPEAKER"):
        return "native"
    for cefr in reversed(LANGUAGE_LEVELS[:-1]):
        if cefr in lvl:
            return cefr
    return ""


def language_rank(level: Optional[str]) -> int:
    """A1=1 … C2=6, native=7, unknown/none=0."""
    lvl = language_level(level) if level else ""
    return LANGUAGE_LEVELS.index(lvl) + 1 if lvl in LANGUAGE_LEVELS else 0


def clean_profile(data: Dict[str, Any]) -> Dict[str, Any]:
    d = data or {}
    links = d.get("links") or {}
    out: Dict[str, Any] = {
        "name": _s(d.get("name"), 120), "headline": _s(d.get("headline"), 200), "location": _s(d.get("location"), 120),
        "summary": _text(d.get("summary"), 2000),
        "links": {k: _s(links.get(k), 300) for k in ("github", "linkedin", "portfolio")},
        "target_categories": [c for c in _list(d.get("target_categories")) if c in config.CATEGORIES and c != "other"],
        "target_seniorities": [c for c in _list(d.get("target_seniorities")) if c in config.SENIORITIES],
        "target_job_types": [c for c in _list(d.get("target_job_types")) if c in config.JOB_TYPES],
        "preferred_locations": _s(d.get("preferred_locations"), 200),
        "weekly_hours": _int(d.get("weekly_hours"), 1, 80, 10),
        "timeline_months": _int(d.get("timeline_months"), 1, 36, 6),
    }
    out["education"] = [
        {"degree": _s(e.get("degree"), 120), "field": _s(e.get("field"), 120), "institution": _s(e.get("institution"), 160),
         "start": _s(e.get("start"), 20), "end": _s(e.get("end"), 20),
         "status": e.get("status") if e.get("status") in EDU_STATUS else "completed"}
        for e in d.get("education") or [] if isinstance(e, dict) and (e.get("degree") or e.get("institution"))
    ]
    out["experience"] = [
        {"title": _s(e.get("title"), 120), "company": _s(e.get("company"), 120), "start": _s(e.get("start"), 20),
         "end": _s(e.get("end"), 20), "current": bool(e.get("current")), "description": _text(e.get("description"), 3000),
         "technologies": _list(e.get("technologies"))}
        for e in d.get("experience") or [] if isinstance(e, dict) and (e.get("title") or e.get("company"))
    ]
    skills, seen = [], set()
    for sk in d.get("skills") or []:
        if not isinstance(sk, dict) or not _s(sk.get("name")):
            continue
        name = ai.canonical_skill(_s(sk.get("name"), 60))
        if name.lower() in seen:
            continue
        seen.add(name.lower())
        years = sk.get("years")
        try:
            years = round(max(0.0, min(50.0, float(years))), 1) if years not in (None, "") else None
        except (TypeError, ValueError):
            years = None
        skills.append({"name": name, "level": _int(sk.get("level"), 1, 4, 2), "years": years})
    out["skills"] = skills
    langs, seen = [], set()
    for l in d.get("languages") or []:
        if not isinstance(l, dict) or not _s(l.get("language")):
            continue
        lang = _s(l.get("language"), 40).title()
        if lang.lower() in seen:
            continue
        seen.add(lang.lower())
        langs.append({"language": lang, "level": language_level(l.get("level"))})
    out["languages"] = langs
    out["certifications"] = [
        {"name": _s(c.get("name"), 160), "issuer": _s(c.get("issuer"), 120), "year": _s(c.get("year"), 10)}
        for c in d.get("certifications") or [] if isinstance(c, dict) and _s(c.get("name"))
    ]
    out["projects"] = [
        {"name": _s(p.get("name"), 120), "description": _text(p.get("description"), 1500),
         "technologies": _list(p.get("technologies")), "url": _s(p.get("url"), 300)}
        for p in d.get("projects") or [] if isinstance(p, dict) and _s(p.get("name"))
    ]
    return out


def merge_profiles(existing: Dict[str, Any], incoming: Dict[str, Any]) -> Dict[str, Any]:
    """Fill the existing profile with CV data without overwriting what the user already entered."""
    a, b = clean_profile(existing), clean_profile(incoming)
    out = dict(a)
    for key in ("name", "headline", "location", "summary"):
        out[key] = a[key] or b[key]
    out["links"] = {k: a["links"][k] or b["links"][k] for k in a["links"]}

    def merge(key: str, ident):
        have = {ident(x) for x in a[key]}
        out[key] = a[key] + [x for x in b[key] if ident(x) not in have]

    merge("education", lambda e: (e["degree"].lower(), e["institution"].lower()))
    merge("experience", lambda e: (e["title"].lower(), e["company"].lower()))
    merge("skills", lambda s: s["name"].lower())
    merge("languages", lambda l: l["language"].lower())
    merge("certifications", lambda c: c["name"].lower())
    merge("projects", lambda p: p["name"].lower())
    return out


# ---------------------------------------------------------------- storage

def ensure_table(conn: sqlite3.Connection) -> None:
    conn.executescript(PROFILE_SCHEMA_SQL)


def load(conn: sqlite3.Connection) -> Tuple[Dict[str, Any], Optional[str]]:
    ensure_table(conn)
    row = conn.execute("SELECT data, cv_filename FROM profile WHERE id=1").fetchone()
    if not row:
        return json.loads(json.dumps(EMPTY_PROFILE)), None
    return clean_profile(json.loads(row["data"])), row["cv_filename"]


def save(conn: sqlite3.Connection, data: Dict[str, Any]) -> Dict[str, Any]:
    ensure_table(conn)
    cleaned = clean_profile(data)
    conn.execute(
        "INSERT INTO profile(id, data, updated_at) VALUES (1, ?, ?) "
        "ON CONFLICT(id) DO UPDATE SET data=excluded.data, updated_at=excluded.updated_at",
        (json.dumps(cleaned), db.now_iso()),
    )
    return cleaned


def save_cv(conn: sqlite3.Connection, filename: str, content: bytes, text: str) -> None:
    ensure_table(conn)
    cv_dir = config.DATA_DIR / "cv"
    cv_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", filename)[-120:] or "cv"
    (cv_dir / safe).write_bytes(content)
    conn.execute(
        "INSERT INTO profile(id, data, cv_filename, cv_text, updated_at) VALUES (1, ?, ?, ?, ?) "
        "ON CONFLICT(id) DO UPDATE SET cv_filename=excluded.cv_filename, cv_text=excluded.cv_text, updated_at=excluded.updated_at",
        (json.dumps(EMPTY_PROFILE), safe, text, db.now_iso()),
    )


def is_filled(p: Dict[str, Any]) -> bool:
    return bool(p.get("skills") or p.get("experience") or p.get("education"))


# ---------------------------------------------------------------- CV reading

class CVError(Exception):
    pass


def extract_cv_text(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    if len(content) > MAX_CV_BYTES:
        raise CVError("The file is larger than 8 MB.")
    if not name.endswith(CV_EXTENSIONS):
        raise CVError("Upload a PDF, Word (.docx) or text file.")
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx"):
            import docx

            document = docx.Document(io.BytesIO(content))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = content.decode("utf-8", errors="replace")
    except CVError:
        raise
    except Exception as exc:
        raise CVError(f"Couldn't read the file: {exc}") from exc
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 80:
        raise CVError("No text found in the file. If it's a scanned PDF, export your CV as a text-based PDF or .docx.")
    return text[:40000]


CV_TOOL_NAME = "record_cv_profile"
CV_SCHEMA = {
    "type": "object",
    "properties": {
        "name": {"type": "string"},
        "headline": {"type": "string", "description": "Short professional headline, e.g. 'Junior Software Engineer (.NET)'."},
        "location": {"type": "string"},
        "summary": {"type": "string", "description": "2-3 sentence professional summary in English, based only on the CV."},
        "links": {"type": "object", "properties": {"github": {"type": "string"}, "linkedin": {"type": "string"}, "portfolio": {"type": "string"}},
                  "required": ["github", "linkedin", "portfolio"]},
        "education": {"type": "array", "items": {"type": "object", "properties": {
            "degree": {"type": "string"}, "field": {"type": "string"}, "institution": {"type": "string"},
            "start": {"type": "string", "description": "YYYY or YYYY-MM, empty if unknown"}, "end": {"type": "string"},
            "status": {"type": "string", "enum": EDU_STATUS}},
            "required": ["degree", "field", "institution", "start", "end", "status"]}},
        "experience": {"type": "array", "description": "Jobs and internships, most recent first.", "items": {"type": "object", "properties": {
            "title": {"type": "string"}, "company": {"type": "string"}, "start": {"type": "string"}, "end": {"type": "string"},
            "current": {"type": "boolean"}, "description": {"type": "string", "description": "What they did, 1-4 sentences."},
            "technologies": {"type": "array", "items": {"type": "string"}}},
            "required": ["title", "company", "start", "end", "current", "description", "technologies"]}},
        "skills": {"type": "array", "items": {"type": "object", "properties": {
            "name": {"type": "string", "description": "Short canonical English name, e.g. 'C#', '.NET', 'SQL', 'Docker'."},
            "level": {"type": "integer", "description": "1 beginner (coursework/brief use), 2 intermediate (used in projects or <2 years at work), 3 advanced (main tool at work 2+ years), 4 expert."},
            "years": {"type": ["number", "null"], "description": "Approximate years of hands-on use, null if unclear."}},
            "required": ["name", "level", "years"]}},
        "languages": {"type": "array", "items": {"type": "object", "properties": {
            "language": {"type": "string"}, "level": {"type": "string", "description": "CEFR A1-C2, 'native', or '' if not stated."}},
            "required": ["language", "level"]}},
        "certifications": {"type": "array", "items": {"type": "object", "properties": {
            "name": {"type": "string"}, "issuer": {"type": "string"}, "year": {"type": "string"}}, "required": ["name", "issuer", "year"]}},
        "projects": {"type": "array", "items": {"type": "object", "properties": {
            "name": {"type": "string"}, "description": {"type": "string"}, "technologies": {"type": "array", "items": {"type": "string"}},
            "url": {"type": "string"}}, "required": ["name", "description", "technologies", "url"]}},
    },
    "required": ["name", "headline", "location", "summary", "links", "education", "experience", "skills", "languages", "certifications", "projects"],
}
CV_SYSTEM = """You turn a CV/résumé into structured profile data for a job-search tool.
- Extract only what the CV states; use "" or [] when something is missing. Never invent employers, dates or degrees.
- Skills: include every technical skill, tool, framework and methodology mentioned (from skill lists AND from job/project descriptions), using short canonical English names. Estimate the level from how and how long it was used.
- Human languages go in "languages", not "skills".
- Translate descriptions to English if the CV is in another language."""


async def parse_cv_with_ai(analyzer: "ai.BaseAnalyzer", cv_text: str) -> Dict[str, Any]:
    raw = await analyzer.structured(CV_TOOL_NAME, "Record the structured profile extracted from the CV.", CV_SCHEMA,
                                    CV_SYSTEM, "CV text:\n\n" + cv_text, max_tokens=6000)
    return clean_profile(raw)


# ---------------------------------------------------------------- helpers used by matching/planning

def skill_levels(p: Dict[str, Any]) -> Dict[str, int]:
    """lowercase canonical skill -> level, including technologies listed under experience/projects (as level 1-2)."""
    levels: Dict[str, int] = {}
    for s in p.get("skills", []):
        levels[s["name"].lower()] = max(levels.get(s["name"].lower(), 0), int(s.get("level") or 2))
    for e in p.get("experience", []):
        for t in e.get("technologies", []):
            k = ai.canonical_skill(t).lower()
            levels[k] = max(levels.get(k, 0), 2)
    for pr in p.get("projects", []):
        for t in pr.get("technologies", []):
            k = ai.canonical_skill(t).lower()
            levels[k] = max(levels.get(k, 0), 1)
    return levels


def language_ranks(p: Dict[str, Any]) -> Dict[str, int]:
    return {l["language"].lower(): language_rank(l.get("level")) for l in p.get("languages", [])}
