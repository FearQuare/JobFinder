"""Profile, CV import, gap analysis, job matching and the learning plan."""
import io
import json

import pytest
from fastapi.testclient import TestClient

from app import ai, db, learning_resources as lr, planner
from app import profile as prof
from tests import test_pipeline as tp
from tests.test_ai_providers import openai_response

PROFILE = {
    "name": "Test User", "headline": "Junior Software Engineer (.NET)", "location": "Vienna",
    "target_categories": ["devops_cloud"], "target_seniorities": ["junior"], "target_job_types": ["full_time"],
    "weekly_hours": 10, "timeline_months": 6,
    "skills": [{"name": "C#", "level": 3, "years": 2}, {"name": "docker", "level": 3}, {"name": "Linux", "level": 1},
               {"name": "k8s", "level": 1}],
    "languages": [{"language": "english", "level": "C1"}, {"language": "Turkish", "level": "native"}],
    "experience": [{"title": "Junior Software Engineer", "company": "Kurulum", "technologies": "SQL, ELK", "current": True}],
    "education": [{"degree": "BSc", "field": "Software Engineering", "institution": "IUE", "status": "completed"}],
}

FAKE_PLAN = {
    "summary": "You have a solid .NET base; the main gaps are Kubernetes and German.",
    "strengths": ["Docker", "C#"],
    "phases": [{"title": "Foundations", "goal": "Comfortable on Linux"}, {"title": "Containers & cloud", "goal": "Deploy to Kubernetes"}],
    "items": [
        {"topic": "Learn German to B1", "skill_key": "German", "kind": "language", "priority": "high",
         "priority_reason": "German B2 is asked for in 100% of target postings.", "phase": 1, "order": 1,
         "current_level": "none", "target_level": "B1", "estimated_hours": 200, "how": "Daily course.", "practice_project": "",
         "resources": [{"title": "Made-up German school", "url": "https://totally-fake-german.example/course", "type": "course", "free": False}]},
        {"topic": "Learn Kubernetes basics", "skill_key": "Kubernetes", "kind": "skill", "priority": "high",
         "priority_reason": "In 100% of junior DevOps postings.", "phase": 2, "order": 3,
         "current_level": "Beginner", "target_level": "Intermediate", "estimated_hours": 40, "how": "Tutorials.",
         "practice_project": "Deploy an app to kind.",
         "resources": [{"title": "Official tutorials", "url": "https://kubernetes.io/docs/tutorials/", "type": "docs", "free": True},
                       {"title": "Bad roadmap", "url": "https://roadmap.sh/not-a-real-roadmap", "type": "roadmap", "free": True}]},
        {"topic": "Linux command line", "skill_key": "Linux", "kind": "skill", "priority": "medium",
         "priority_reason": "40% of postings.", "phase": 1, "order": 2, "current_level": "Beginner", "target_level": "Intermediate",
         "estimated_hours": 25, "how": "Linux Journey.", "practice_project": "", "resources": []},
        {"topic": "Learn Excel", "skill_key": "", "kind": "skill", "priority": "low", "priority_reason": "Rarely asked.",
         "phase": 9, "order": "x", "current_level": "", "target_level": "", "estimated_hours": "5", "how": "", "practice_project": "",
         "resources": []},
    ],
    "job_search_tips": ["Mention German progress on your CV."],
}


class PlanFake(tp.FakeAnalyzer):
    def __init__(self, *a, **k):
        super().__init__()
        self.prompts = []

    async def structured(self, name, description, schema, system, prompt, max_tokens=2000):
        self.prompts.append((name, prompt))
        if name == planner.PLAN_TOOL_NAME:
            return json.loads(json.dumps(FAKE_PLAN))
        if name == prof.CV_TOOL_NAME:
            return {"name": "Elmira CV", "headline": "Software Engineer", "location": "Izmir", "summary": "Dev.",
                    "links": {"github": "https://github.com/x", "linkedin": "", "portfolio": ""},
                    "education": [{"degree": "BSc", "field": "Software Engineering", "institution": "IUE", "start": "2020", "end": "2024", "status": "completed"}],
                    "experience": [{"title": "Intern", "company": "ASELSAN", "start": "2023", "end": "2023", "current": False,
                                    "description": "Web app testing", "technologies": ["React", "Selenium"]}],
                    "skills": [{"name": "Selenium", "level": 2, "years": 1}, {"name": "C#", "level": 1, "years": None}],
                    "languages": [{"language": "German", "level": "A1"}],
                    "certifications": [], "projects": []}
        raise AssertionError(name)


# ---------------------------------------------------------------- profile

def test_clean_profile_normalises_and_drops_junk():
    p = prof.clean_profile({**PROFILE, "target_categories": ["devops_cloud", "other", "nonsense"],
                            "skills": PROFILE["skills"] + [{"name": ""}, {"name": "Docker", "level": 9}],
                            "weekly_hours": "500"})
    assert p["target_categories"] == ["devops_cloud"]
    assert [s["name"] for s in p["skills"]] == ["C#", "docker", "Linux", "Kubernetes"]
    assert p["weekly_hours"] == 80
    assert p["languages"] == [{"language": "English", "level": "C1"}, {"language": "Turkish", "level": "native"}]
    assert p["experience"][0]["technologies"] == ["SQL", "ELK"]
    assert prof.language_rank("B2") == 4 and prof.language_rank("native") == 7 and prof.language_rank("") == 0
    assert prof.language_level("Deutsch – sehr gut (C1)") == "C1"


def test_merge_keeps_user_entries_and_adds_new():
    merged = prof.merge_profiles(PROFILE, {"name": "Other", "skills": [{"name": "C#", "level": 1}, {"name": "Selenium", "level": 2}],
                                           "languages": [{"language": "German", "level": "A1"}]})
    assert merged["name"] == "Test User"
    skills = {s["name"]: s["level"] for s in merged["skills"]}
    assert skills["C#"] == 3 and skills["Selenium"] == 2
    assert [l["language"] for l in merged["languages"]] == ["English", "Turkish", "German"]
    assert merged["target_categories"] == ["devops_cloud"]


def _docx_bytes(lines):
    import docx

    d = docx.Document()
    for line in lines:
        d.add_paragraph(line)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Skills", "C#, SQL, Docker"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf_bytes(text):
    """Minimal valid single-page PDF with one line of Helvetica text."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def test_extract_cv_text_formats():
    long_line = "Junior Software Engineer at Kurulum Cognitive Services working with C# .NET SQL and ELK stack"
    assert "Kurulum" in prof.extract_cv_text("cv.pdf", _pdf_bytes(long_line))
    text = prof.extract_cv_text("CV.docx", _docx_bytes([long_line, "Education: BSc Software Engineering"]))
    assert "Kurulum" in text and "Skills | C#, SQL, Docker" in text
    assert "Kurulum" in prof.extract_cv_text("cv.txt", long_line.encode())
    with pytest.raises(prof.CVError, match="PDF, Word"):
        prof.extract_cv_text("cv.png", b"x" * 100)
    with pytest.raises(prof.CVError, match="No text found"):
        prof.extract_cv_text("cv.txt", b"too short")


# ---------------------------------------------------------------- market, gaps, matching

@pytest.fixture
def market_db(tmp_data, monkeypatch):
    karriere = tp.FakeSource("karriere", tp._karriere_pages(), details={"k1": "Kubernetes", "k2": "Docker"})
    linkedin = tp.FakeSource("linkedin", tp._linkedin_pages())
    state = tp._run_collection(monkeypatch, karriere, linkedin)
    assert state["status"] == "done"
    return tmp_data


def test_gap_analysis_urgency(market_db):
    p = prof.clean_profile(PROFILE)
    with db.session() as conn:
        market = planner.market_for_profile(conn, p, 14)
    # only 1 junior DevOps job (< 5) → falls back to all DevOps levels (3 jobs)
    assert market["groups"] == [{"category": "devops_cloud", "seniority": None, "job_count": 3}]
    kube = next(s for s in market["skills"] if s["skill"] == "Kubernetes")
    assert kube["pct"] == 100 and kube["required_pct"] == 67
    gaps = planner.gap_analysis(p, market)
    by = {g["skill"]: g for g in gaps["skills"]}
    assert by["Docker"]["status"] == "strong" and by["Docker"]["suggested_priority"] is None
    assert by["Kubernetes"]["status"] == "partial" and by["Kubernetes"]["suggested_priority"] == "high"
    assert by["Terraform"]["status"] == "missing"
    german = gaps["languages"][0]
    assert german["language"] == "German" and german["typical_level"] == "B2" and german["suggested_priority"] == "high"
    assert 0 < gaps["match_score"] < 100
    assert "Docker" in gaps["strengths"]


def test_job_match_scores(market_db):
    p = prof.clean_profile(PROFILE)
    with db.session() as conn:
        ids = [r[0] for r in conn.execute("SELECT id FROM jobs WHERE category='devops_cloud'")]
        titles = {r["id"]: r["title"] for r in conn.execute("SELECT id, title FROM jobs")}
        scores = planner.job_match_scores(conn, p, ids)
    junior = next(s for i, s in scores.items() if titles[i].startswith("Junior"))
    # requires Docker (have 3) + Linux (have 1 → half), Kubernetes nice (have 1 → half); German B2 required → ×0.6
    assert junior["language_gap"] == "German B2"
    assert junior["missing"] == ["Linux"]
    assert junior["score"] == round(100 * (2 + 1 + 0.5) / 5 * 0.6)
    p2 = prof.clean_profile({**PROFILE, "languages": [{"language": "German", "level": "C1"}]})
    with db.session() as conn:
        assert planner.job_match_scores(conn, p2, ids)[next(i for i in ids if titles[i].startswith("Junior"))]["language_gap"] is None


# ---------------------------------------------------------------- plan

def test_postprocess_plan_validates_links_and_builds_timeline(market_db):
    p = prof.clean_profile(PROFILE)
    with db.session() as conn:
        market = planner.market_for_profile(conn, p, 14)
    gaps = planner.gap_analysis(p, market)
    plan = planner.postprocess_plan(json.loads(json.dumps(FAKE_PLAN)), p, market, gaps)
    items = {it["skill_key"] or it["topic"]: it for it in plan["items"]}
    german = items["German"]
    urls = [r["url"] for r in german["resources"]]
    assert urls[:3] == [r["url"] for r in lr.RESOURCES["german"]]          # catalog first, verified
    assert all(r["verified"] for r in german["resources"][:3])
    fake = german["resources"][3]
    assert fake["url"].startswith("https://duckduckgo.com/?q=") and not fake["verified"]   # invented URL replaced
    kube = items["Kubernetes"]
    assert kube["roadmap"] == {"title": "roadmap.sh: Kubernetes", "url": "https://roadmap.sh/kubernetes"}
    assert kube["market"]["pct"] == 100
    bad = [r for r in kube["resources"] if r["title"] == "Bad roadmap"][0]
    assert not bad["verified"] and "duckduckgo" in bad["url"]
    excel = items["Learn Excel"]
    assert excel["phase"] == 2 and excel["estimated_hours"] == 5    # out-of-range phase clamped, strings coerced
    assert [it["topic"] for it in plan["items"]][:2] == ["Learn German to B1", "Linux command line"]  # phase, then order
    assert [it["order"] for it in plan["items"]] == [1, 2, 3, 4]
    # languages run in parallel: phase 1 hours only count Linux (25 h → 3 weeks at 10 h/week)
    assert plan["phases"][0]["hours"] == 25 and plan["phases"][0]["weeks"] == 3
    assert plan["counts"] == {"high": 2, "medium": 1, "low": 1}
    assert {r["slug"] for r in plan["role_roadmaps"]} == {"devops-beginner", "devops"}


def test_plan_and_cv_schemas_are_openai_strict_compatible():
    for schema in (planner.PLAN_SCHEMA, prof.CV_SCHEMA):
        strict = ai.strict_json_schema(schema)

        def walk(node):
            if isinstance(node, dict):
                if node.get("type") == "object" and "properties" in node:
                    assert node["additionalProperties"] is False
                    assert set(node["required"]) == set(node["properties"])
                    assert set(schema_required(node)) == set(node["properties"])
                for v in node.values():
                    walk(v)
            elif isinstance(node, list):
                for v in node:
                    walk(v)

        walk(strict)


def schema_required(node):
    return node["required"]


def test_openai_structured_plan_roundtrip():
    import httpx

    seen = {}

    def handler(request):
        seen["body"] = json.loads(request.content)
        return httpx.Response(200, json=openai_response(json.dumps(FAKE_PLAN)))

    analyzer = ai.OpenAIAnalyzer("sk", "gpt-6-luna", http_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)), max_retries=0)
    import asyncio

    out = asyncio.run(analyzer.structured(planner.PLAN_TOOL_NAME, "d", planner.PLAN_SCHEMA, planner.PLAN_SYSTEM, "prompt", max_tokens=12000))
    assert out["items"][0]["topic"] == "Learn German to B1"
    assert seen["body"]["text"]["format"]["name"] == planner.PLAN_TOOL_NAME
    assert seen["body"]["max_output_tokens"] == 16000


def test_profile_and_plan_pages(market_db, monkeypatch):
    from app.main import app

    fake = PlanFake()
    monkeypatch.setattr(ai, "make_analyzer", lambda settings, concurrency=None, purpose="jobs": fake)
    with TestClient(app) as client:
        r = client.get("/profile")
        assert r.status_code == 200 and "Import from your CV" in r.text
        r = client.get("/plan")
        assert "Start with your profile" in r.text

        # CV import fills an empty profile (not saved until the user saves)
        r = client.post("/api/profile/cv", files={"file": ("cv.docx", _docx_bytes(["Intern at ASELSAN testing web apps with Selenium and React"] * 3))})
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["found"]["experience"] == 1 and body["profile"]["name"] == "Elmira CV"
        assert "ASELSAN" in fake.prompts[-1][1]

        r = client.post("/api/profile", json=PROFILE)
        assert r.status_code == 200 and r.json()["profile"]["skills"][3]["name"] == "Kubernetes"

        r = client.get("/plan")
        assert "Quick check" in r.text and "German" in r.text

        r = client.post("/api/plan/generate")
        assert r.status_code == 200, r.text
        plan_id = r.json()["id"]
        prompt = fake.prompts[-1][1]
        assert "SKILL GAPS WITH COMPUTED URGENCY" in prompt and "Kubernetes" in prompt and "learngerman.dw.com" in prompt

        r = client.get(f"/plan?id={plan_id}")
        assert r.status_code == 200
        assert "High priority" in r.text and "Learn German to B1" in r.text and "roadmap.sh: DevOps" in r.text
        assert "https://totally-fake-german.example" not in r.text

        r = client.post(f"/api/plan/{plan_id}/items/i2", json={"status": "done"})
        assert r.status_code == 200 and r.json()["progress_pct"] > 0
        assert client.post(f"/api/plan/{plan_id}/items/i2", json={"status": "bogus"}).status_code == 400

        r = client.get("/?sort=match")
        assert r.status_code == 200 and "match-pill" in r.text and "Best match for me" in r.text
        assert "needs German B2" in r.text


def test_cv_import_without_key(tmp_data):
    from app.main import app

    with TestClient(app) as client:
        r = client.post("/api/profile/cv", files={"file": ("cv.txt", b"Software engineer with C# and SQL experience " * 5)})
        assert r.status_code == 400 and "API key" in r.json()["error"]
        with db.session() as conn:
            assert conn.execute("SELECT cv_filename FROM profile").fetchone()[0] == "cv.txt"
        r = client.post("/api/plan/generate")
        assert r.status_code == 400 and "profile" in r.json()["error"]


# ---------------------------------------------------------------- certifications & roadmaps

def test_catalog_integrity():
    for cid, c in lr.CERTIFICATIONS.items():
        assert c["url"].startswith("https://") and c["level"] in lr.LEVEL_ORDER, cid
        assert c["url"] in lr.KNOWN_URLS
    assert "az-204" not in lr.CERTIFICATIONS  # retired
    for cat, ids in lr.ROLE_PATHS.items():
        assert cat in lr.ROLE_ROADMAPS and all(i in lr.LEARNING_PATHS for i in ids)
    road = lr.roadmap_catalog(["devops_cloud"], ["junior"])
    ids = [r["id"] for r in road]
    assert ids[:2] == ["roadmapsh-devops-beginner", "roadmapsh-devops"] and "cloud-resume-challenge" in ids
    assert len({r["provider"] for r in road}) >= 5
    # a certificate the AI returns as a resource URL counts as verified
    assert lr.safe_resource({"title": "CKA", "url": lr.CERTIFICATIONS["cka"]["url"]}, "k8s")["verified"]


def test_cert_mentions_matching():
    mentioned = {"cka": 3, "certified kubernetes administrator (cka)": 1, "az-104": 2, "istqb foundation level": 4}
    assert lr.cert_mentions(lr.CERTIFICATIONS["cka"], mentioned) == 4
    assert lr.cert_mentions(lr.CERTIFICATIONS["az-104"], mentioned) == 2
    assert lr.cert_mentions(lr.CERTIFICATIONS["istqb-ctfl"], mentioned) == 4
    assert lr.cert_mentions(lr.CERTIFICATIONS["ckad"], mentioned) == 0


def test_plan_options_schema_and_postprocess_certs(market_db):
    p = prof.clean_profile({**PROFILE, "certifications": [{"name": "Microsoft Azure Fundamentals AZ-900", "issuer": "Microsoft"}]})
    with db.session() as conn:
        jid = conn.execute("SELECT id FROM jobs WHERE category='devops_cloud' LIMIT 1").fetchone()[0]
        conn.execute("INSERT INTO job_skills(job_id, skill, skill_type, required) VALUES (?, 'CKA', 'certification', 0)", (jid,))
        market = planner.market_for_profile(conn, p, 14)
    assert market["cert_mentions"] == {"cka": 1}
    gaps = planner.gap_analysis(p, market)
    options = planner.plan_options(p, market, gaps)
    cert_ids = [c["id"] for c in options["certs"]]
    assert "az-900" not in cert_ids                     # already held
    assert "cka" in cert_ids and "osd" in cert_ids      # role cert + German (a gap) certificate
    assert next(c for c in options["certs"] if c["id"] == "cka")["mentions"] == 1
    schema = planner.plan_schema(options)
    assert schema["properties"]["certifications"]["items"]["properties"]["cert_id"]["enum"] == cert_ids
    assert "recommended_roadmaps" in schema["required"]
    strict = ai.strict_json_schema(schema)          # still valid for OpenAI strict mode / Claude structured outputs
    assert strict["properties"]["certifications"]["items"]["additionalProperties"] is False
    prompt = planner.build_plan_prompt(p, market, gaps, options)
    assert "CERTIFICATION OPTIONS" in prompt and '"mentioned_in_postings": 1' in prompt and "cloud-resume-challenge" in prompt

    raw = json.loads(json.dumps(FAKE_PLAN))
    raw["certifications"] = [
        {"cert_id": "cka", "priority": "medium", "phase": 2, "why": "Named in postings.", "prep_hours": 60},
        {"cert_id": "osd", "priority": "high", "phase": 1, "why": "Proof of German.", "prep_hours": "20"},
        {"cert_id": "made-up-cert", "priority": "high", "phase": 1, "why": "x", "prep_hours": 1},
        {"cert_id": "cka", "priority": "low", "phase": 2, "why": "dupe", "prep_hours": 1},
    ]
    raw["recommended_roadmaps"] = [{"roadmap_id": "cloud-resume-challenge", "how_to_use": "Portfolio."},
                                   {"roadmap_id": "nope", "how_to_use": "x"}]
    plan = planner.postprocess_plan(raw, p, market, gaps, options)
    assert [c["id"] for c in plan["certifications"]] == ["osd", "cka"]     # high first, invalid + dupes dropped
    cka = plan["certifications"][1]
    assert cka["url"] == lr.CERTIFICATIONS["cka"]["url"] and cka["mentions"] == 1 and cka["prep_hours"] == 60
    assert plan["phases"][1]["certifications"] == ["Certified Kubernetes Administrator (CKA)"]
    assert [r["id"] for r in plan["recommended_roadmaps"]] == ["cloud-resume-challenge"]
    assert "cloud-resume-challenge" not in [r["id"] for r in plan["other_roadmaps"]]


def test_model_choice_in_ui(market_db, monkeypatch):
    from app.main import app

    used = {}

    class Recording(PlanFake):
        pass

    def fake_make(settings, concurrency=None, purpose="jobs"):
        used["purpose"] = purpose
        used["choice"] = ai.planning_choice(settings)
        used["effort"] = settings.get("planning_effort")
        f = Recording()
        f.model = used["choice"]["model"] if used["choice"] else "x"
        return f

    monkeypatch.setattr(ai, "make_analyzer", fake_make)
    with TestClient(app) as client:
        client.post("/api/profile", json=PROFILE)
        r = client.get("/plan")
        assert 'id="planning-model"' in r.text and "Claude Opus 5.5 ★" in r.text and "no key" in r.text  # OpenAI has no key
        r = client.post("/api/plan/generate", json={"model": "anthropic:claude-fable-5-1", "effort": "medium"})
        assert r.status_code == 200, r.text
        assert used == {"purpose": "planning", "choice": {"provider": "anthropic", "model": "claude-fable-5-1"}, "effort": "medium"}
        r = client.get(f"/plan?id={r.json()['id']}")
        assert "claude-fable-5-1" in r.text
        # the choice is remembered, and a model without a key is refused clearly
        r = client.post("/api/plan/generate", json={"model": "openai:gpt-6-astra"})
        assert r.status_code == 400 and "OpenAI" in r.json()["error"]
        # settings page: custom model
        client.post("/settings/options", data={"planning_model": "custom", "planning_custom_provider": "anthropic",
                                               "planning_custom_model": "claude-opus-5-5-preview", "planning_effort": "low"})
        with db.session() as conn:
            s = db.get_settings(conn)
        assert s["planning_model"] == "anthropic:claude-opus-5-5-preview" and s["planning_effort"] == "low"
        r = client.get("/settings")
        assert "AI for profile analysis" in r.text and "claude-opus-5-5-preview" in r.text
        # CV import uses the planning model too
        client.post("/api/profile/cv", files={"file": ("cv.txt", b"Software engineer with C# and SQL experience at Kurulum " * 4)})
        assert used["purpose"] == "planning" and used["choice"]["model"] == "claude-opus-5-5-preview"
